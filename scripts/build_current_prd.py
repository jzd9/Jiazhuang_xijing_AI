from docx import Document
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_BREAK
from pathlib import Path


OUT = Path('/Users/nd/Desktop/Jiazhuang_AI/PRD_V2.1_AI家居空间预览_当前实现与测试版.docx')

NAVY = '173B3A'
GREEN = '1E6A5C'
MINT = 'EAF4F0'
ORANGE = 'D97832'
LIGHT_ORANGE = 'FFF1E6'
INK = '1D2928'
GRAY = '5F6D69'
LIGHT_GRAY = 'F3F6F5'
BORDER = 'D8E2DF'
RED = 'B54A45'
YELLOW = 'FFF5D8'
WHITE = 'FFFFFF'


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tc_pr.append(shd)
    shd.set(qn('w:fill'), fill)


def set_cell_border(cell, color=BORDER, size='4'):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in('w:tcBorders')
    if borders is None:
        borders = OxmlElement('w:tcBorders')
        tc_pr.append(borders)
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        tag = 'w:' + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn('w:val'), 'single')
        element.set(qn('w:sz'), size)
        element.set(qn('w:color'), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement('w:tblHeader')
    tbl_header.set(qn('w:val'), 'true')
    tr_pr.append(tbl_header)


def set_cell_text(cell, text, bold=False, color=INK, size=9.2):
    cell.text = ''
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = 1.12
    r = p.add_run(str(text))
    r.bold = bold
    r.font.name = 'Hiragino Sans GB'
    r._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'Hiragino Sans GB')
    r.font.size = Pt(size)
    r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_border(cell)


def add_table(doc, headers, rows, widths=None, font_size=9.0):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for i, h in enumerate(headers):
        set_cell_text(hdr.cells[i], h, bold=True, color=WHITE, size=9.2)
        set_cell_shading(hdr.cells[i], NAVY)
        if widths:
            hdr.cells[i].width = Cm(widths[i])
    for row_idx, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value, size=font_size)
            set_cell_shading(cells[i], WHITE if row_idx % 2 == 0 else LIGHT_GRAY)
            if widths:
                cells[i].width = Cm(widths[i])
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_label(doc, text, fill=MINT, color=GREEN):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    cell = table.cell(0, 0)
    set_cell_text(cell, text, bold=True, color=color, size=9.5)
    set_cell_shading(cell, fill)
    set_cell_border(cell, fill, '0')
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_callout(doc, title, body, kind='info'):
    fills = {'info': MINT, 'warning': YELLOW, 'risk': 'FCEBE9', 'note': LIGHT_GRAY}
    colors = {'info': GREEN, 'warning': '8A6500', 'risk': RED, 'note': GRAY}
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    cell.text = ''
    set_cell_shading(cell, fills[kind])
    set_cell_border(cell, fills[kind], '0')
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(title)
    r.bold = True
    r.font.size = Pt(10.5)
    r.font.color.rgb = RGBColor.from_string(colors[kind])
    r._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'Hiragino Sans GB')
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.line_spacing = 1.2
    r2 = p2.add_run(body)
    r2.font.size = Pt(9.5)
    r2.font.color.rgb = RGBColor.from_string(INK)
    r2._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'Hiragino Sans GB')
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_para(doc, text='', style=None, bold_prefix=None, color=None, align=None, keep=False):
    p = doc.add_paragraph(style=style)
    if align:
        p.alignment = align
    if keep:
        p.paragraph_format.keep_with_next = True
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        r1.bold = True
        r2 = p.add_run(text[len(bold_prefix):])
        runs = [r1, r2]
    else:
        runs = [p.add_run(text)]
    if color:
        for r in runs:
            r.font.color.rgb = RGBColor.from_string(color)
    return p


def add_bullets(doc, items, level=0):
    for item in items:
        style = 'List Bullet' if level == 0 else 'List Bullet 2'
        add_para(doc, item, style=style)


def add_numbered(doc, items):
    for item in items:
        add_para(doc, item, style='List Number')


def add_code(doc, code):
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    set_cell_shading(cell, '142321')
    set_cell_border(cell, '142321', '0')
    cell.text = ''
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    r = p.add_run(code)
    r.font.name = 'Menlo'
    r._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'Menlo')
    r.font.size = Pt(7.7)
    r.font.color.rgb = RGBColor.from_string('E8F0EE')
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def page_break(doc):
    doc.add_page_break()


doc = Document()
section = doc.sections[0]
section.top_margin = Cm(1.8)
section.bottom_margin = Cm(1.7)
section.left_margin = Cm(2.0)
section.right_margin = Cm(2.0)

# Theme styles
styles = doc.styles
normal = styles['Normal']
normal.font.name = 'Hiragino Sans GB'
normal._element.rPr.rFonts.set(qn('w:eastAsia'), 'Hiragino Sans GB')
normal.font.size = Pt(10)
normal.font.color.rgb = RGBColor.from_string(INK)
normal.paragraph_format.space_after = Pt(5)
normal.paragraph_format.line_spacing = 1.25

for name, size, color, before, after in [
    ('Title', 31, NAVY, 0, 10),
    ('Subtitle', 15, GREEN, 0, 12),
    ('Heading 1', 20, NAVY, 18, 8),
    ('Heading 2', 14, GREEN, 14, 5),
    ('Heading 3', 11, ORANGE, 10, 4),
]:
    st = styles[name]
    st.font.name = 'Hiragino Sans GB'
    st._element.rPr.rFonts.set(qn('w:eastAsia'), 'Hiragino Sans GB')
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = RGBColor.from_string(color)
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)
    st.paragraph_format.keep_with_next = True

for name in ['List Bullet', 'List Bullet 2', 'List Number']:
    st = styles[name]
    st.font.name = 'Hiragino Sans GB'
    st._element.rPr.rFonts.set(qn('w:eastAsia'), 'Hiragino Sans GB')
    st.font.size = Pt(10)
    st.font.color.rgb = RGBColor.from_string(INK)
    st.paragraph_format.space_after = Pt(3)

caption = styles['Caption']
caption.font.name = 'Hiragino Sans GB'
caption._element.rPr.rFonts.set(qn('w:eastAsia'), 'Hiragino Sans GB')
caption.font.size = Pt(8.5)
caption.font.color.rgb = RGBColor.from_string(GRAY)

# Footer with page number
footer = section.footer
p = footer.paragraphs[0]
p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r = p.add_run('栖境 AI · PRD V2.1  |  ')
r.font.size = Pt(8)
r.font.color.rgb = RGBColor.from_string(GRAY)
fld_char1 = OxmlElement('w:fldChar'); fld_char1.set(qn('w:fldCharType'), 'begin')
instr_text = OxmlElement('w:instrText'); instr_text.set(qn('xml:space'), 'preserve'); instr_text.text = 'PAGE'
fld_char2 = OxmlElement('w:fldChar'); fld_char2.set(qn('w:fldCharType'), 'end')
r._r.append(fld_char1); r._r.append(instr_text); r._r.append(fld_char2)

# Cover
add_label(doc, 'CURRENT STATE / TESTED BASELINE', fill=LIGHT_ORANGE, color=ORANGE)
doc.add_paragraph().paragraph_format.space_after = Pt(36)
p = doc.add_paragraph(style='Title')
p.add_run('AI 家居空间预览')
p2 = doc.add_paragraph(style='Subtitle')
p2.add_run('产品需求文档 PRD V2.1')
p3 = doc.add_paragraph()
p3.paragraph_format.space_after = Pt(20)
r = p3.add_run('当前实现与 V1→V2 测试归档版')
r.font.size = Pt(16)
r.font.bold = True
r.font.color.rgb = RGBColor.from_string(ORANGE)
add_callout(doc, '文档定位', '本文件以 2026-10-09 的 V2_test 可运行基线、当前部署页面、原 PRD V2、12 条测试记录及 V1→V2 修改记录为依据。它描述“现在真实能做什么”，并把未实现或证据不足的内容明确标为“待确认”或“后续规划”。', 'info')
doc.add_paragraph().paragraph_format.space_after = Pt(24)
add_table(doc, ['项目项', '内容'], [
    ['产品名称', '栖境 AI · AI 家居空间预览'],
    ['文档版本', 'V2.1（当前实现与测试归档版）'],
    ['基线版本', 'V2_test'],
    ['基线提交', '7f9f3c03d1ac6b652aa9a8c649168660a821e882'],
    ['当前演示地址', 'https://qijing-ai-home-preview.mawellhihi.chatgpt.site/'],
    ['文档日期', '2026-10-09'],
    ['原始方案', 'PRD_V2_AI家居空间预览.docx（保留不覆盖）'],
], widths=[4.0, 12.0], font_size=9.5)
add_para(doc, '保密与使用说明：本文档用于产品、设计、研发、测试和后续 AI 协作者对齐。测试结论仅覆盖所提供的 12 条样例，不能代表大规模线上质量。', color=GRAY)

page_break(doc)
doc.add_heading('文档说明与证据边界', level=1)
add_para(doc, '本 PRD 的事实优先级依次为：V2_test 当前代码与可运行页面 → 本轮测试记录 → V1→V2 修改记录 → 原 PRD V2 的产品设想。原 PRD 中与当前实现不一致的内容，不继续写成“已完成”。')
add_table(doc, ['材料', '用途', '处理原则'], [
    ['V2_test 前端与 Worker', '确认真实页面、交互、状态、调用链和本地存储', '作为“已实现”判断的最高优先级'],
    ['12 条测试结果（T01–T12）', '确认 V1/V2 表现、回归风险和质量边界', '保留通过/半通过/失败的原始结论'],
    ['V1→V2 修改记录', '还原问题、优先级、已做修复和未完成方向', '与当前代码交叉验证后再归类'],
    ['原 PRD V2', '继承产品背景、目标、概念和业务意图', '冲突项降级为“历史方案”或“待确认”'],
], widths=[4.0, 5.4, 6.6])
add_callout(doc, '状态词定义', '已实现：当前基线中存在且可从代码或页面验证；部分实现：有能力但存在明确限制；待确认：材料不足，无法确认产品决策或线上能力；后续规划：测试记录或原 PRD 提出、当前基线尚未实现。', 'note')

doc.add_heading('执行摘要', level=1)
add_para(doc, '栖境 AI 是一个面向普通家居用户的“聊天式空间编辑”原型。用户上传空间照片和家具商品图，通过点选落点、选择添加或替换、填写可选尺寸并输入自然语言，即可让系统调用阿里云百炼模型，把家具以场景理解式图像编辑的方式融入原空间。')
add_para(doc, '当前版本已从早期“抠图叠加”的拼贴路径升级为“原房间作为主编辑图、商品图作为外观参考”的生成路径；正常链路包含两轮实际生图和一轮视觉质检。前台只展示最终效果，管理员后台可查看 Round 1 与 Round 2 两张过程图、每轮 Prompt、耗时、状态和请求编号。')
add_para(doc, '系统当前仍是单浏览器演示级实现：历史对话、Prompt 配置和任务摘要主要保存在浏览器 localStorage；管理员页没有登录鉴权；长任务、落点精度、语义冲突识别和模糊素材提示仍需继续优化。')
add_table(doc, ['结论项', '当前判断'], [
    ['核心任务是否可完成', '可以。添加、替换、多轮引用、下载、失败重试均已实现。'],
    ['是否真实调用接口', '是。Worker 调用百炼视觉理解、图像编辑和视觉质检接口，并保存 request_id。'],
    ['是否为纯前端贴图', '否。最终图来自服务端图像编辑生成；商品图不再作为浏览器叠加层。'],
    ['测试表现', 'V2：10/12 严格通过，1/12 半通过，1/12 失败。'],
    ['上线成熟度', '演示/验证阶段；尚不满足生产级账号、权限、持久化、超时治理和大规模稳定性要求。'],
], widths=[4.2, 11.8], font_size=9.4)

page_break(doc)
doc.add_heading('1. 产品背景', level=1)
add_para(doc, '消费者在线购买家具时，通常只能看到白底商品图或经过布景的宣传图，很难判断家具放进自己的真实空间后是否合适。核心不确定性包括：尺寸比例、摆放位置、透视、接地、遮挡、光影、材质和整体风格。传统手工合成需要专业工具，直接抠图贴图又容易产生“漂浮、穿模、比例错误、背景失真”等问题。')
add_para(doc, '本产品尝试把这类任务转化为低门槛的聊天式操作：用户提供自己的空间图和商品图，用自然语言表达“放在哪里、添加还是替换、是否调整大小”，系统负责理解空间并生成融合后的预览图。')

doc.add_heading('2. 目标用户', level=1)
add_table(doc, ['用户', '主要需求', '当前支持程度'], [
    ['普通家居消费者', '购买前预览家具在自家空间中的效果', '核心目标用户，已支持'],
    ['租房或局部改造用户', '只改一个局部物件，不重做整个房间', '已支持局部添加/替换'],
    ['软装顾问/门店导购', '快速用客户照片生成沟通稿', '可演示使用；团队协作与客户管理未实现'],
    ['运营/管理员', '查看真实调用、过程图、Prompt 和耗时', '后台审阅已支持；权限与集中存储未实现'],
], widths=[3.3, 7.2, 5.5])

doc.add_heading('3. 典型使用场景', level=1)
add_bullets(doc, [
    '添加：在空房、桌面、柜面或局部空间中加入沙发、落地灯、椅子、花瓶等物体。场景不限定为“客厅/卧室”等房间类型。',
    '替换：点选现有家具，用商品图中的新家具替换它，同时尽量保留背景与附近独立小物品。',
    '连续调整：引用刚生成的效果图，继续说“整体小一点”“向右移动一些”“位置不变，只调整比例”。',
    '历史引用：在同一对话内引用过去某张效果图作为下一轮主编辑图。',
    '异常恢复：接口失败后直接重新生成，不重新上传素材或输入指令。',
    '运营审阅：管理员查看每轮真实过程图、Prompt 版本、耗时、状态和 request_id。',
])

doc.add_heading('4. 产品目标与原则', level=1)
add_table(doc, ['目标/原则', '产品解释', '当前状态'], [
    ['降低预览门槛', '以聊天、上传、点选代替专业修图流程', '已实现'],
    ['结果像真实场景编辑', '家具应融入空间，不像平面贴图', '已实现生成路径；质量仍受模型能力影响'],
    ['保留原房间', '结构、拍摄角度、纹理、曝光和颗粒尽量不变', 'Prompt 与输入链路已约束'],
    ['让用户可连续迭代', '引用结果继续修改，保持上下文与版本', '已实现线性引用；正式版本树待确认'],
    ['过程可审计', '后台可见真实轮次、耗时和 Prompt', '已实现本地后台审阅'],
    ['错误可恢复', '失败不丢素材，提供直接重试', '已实现；浏览器刷新后的大图素材恢复受存储限制'],
], widths=[3.2, 8.0, 4.8])

doc.add_heading('4.1 本阶段非目标', level=2)
add_bullets(doc, [
    '不是完整室内设计工具，不提供户型建模、施工图、预算或材料清单。',
    '不承诺厘米级测量精度，当前尺寸用于生成约束和比例参考。',
    '不提供多人协作、账号体系、云端跨设备历史同步。',
    '不提供商品购买、库存、价格或电商交易闭环。',
    '不保证任何输入都能一次得到可用结果；当前允许最多一次自动纠错。',
])

page_break(doc)
doc.add_heading('5. 当前产品范围概览', level=1)
add_table(doc, ['模块', '已实现能力', '关键限制'], [
    ['对话与历史', '新建、切换、搜索、自动标题、本地保存、删除确认', '仅当前浏览器；最多约 12 个对话'],
    ['素材输入', '空间图、商品图上传；尺寸与元数据记录', '空间图 10MB；商品图会在前端缩放为参考图'],
    ['编辑意图', '添加/替换、文字指令、尺寸、点选落点', '无画框/蒙版；点选表示底部中心接触点'],
    ['冲突与落点', '左右冲突确认；轻微偏差自动修正；明显错误确认', '推荐位置只显示坐标，未可视化标记'],
    ['生成', 'Round 0 解析、Round 1 基础、Round 2 综合修正、Round 3 QC', '长任务与模型服务稳定性仍有风险'],
    ['结果', '查看、引用、下载、重新生成', '没有正式分支树或版本对比'],
    ['管理员后台', '任务列表、过程图、Prompt/耗时/状态/request_id', '无账号权限；数据主要本地保存'],
    ['Prompt 管理', 'Round 0–3 分别查看、编辑、保存、取消、恢复默认', '保存即生效；无草稿/发布流程'],
], widths=[3.0, 8.0, 5.0], font_size=8.8)

doc.add_heading('6. 用户端功能说明', level=1)
doc.add_heading('6.1 历史对话', level=2)
add_bullets(doc, [
    '左侧列表每一项代表一段独立生图对话，显示简短标题和结果图数量。',
    '标题从该对话第一轮指令自动生成，不按固定房间类型命名。',
    '点击“新对话”会归档当前对话并进入空状态；不会把上一主题消息带入新对话。',
    '删除入口默认弱化，点击后需二次确认；若有生成任务进行中，则禁止删除。',
    '删除当前对话后自动切换到最近可用对话；没有剩余对话时进入新建空状态。',
    '删除范围包括本地消息、结果、版本引用关系和该对话任务摘要；不影响其他对话。',
])
add_callout(doc, '数据事实', '历史对话删除目前是前端本地删除，不代表后端百炼侧任务或远程生成文件被删除。是否需要服务端级数据删除与保留周期，待确认。', 'warning')

doc.add_heading('6.2 素材上传与原图保护', level=2)
add_bullets(doc, [
    '空间原图支持 PNG、JPEG、WebP，单文件上限 10MB。系统记录 original_width、original_height、aspect_ratio、文件名和大小。',
    '生图基图优先使用用户最初上传的原始图片数据，不使用页面预览截图。',
    '前端预览可在固定容器中自适应显示，但展示尺寸与真实生图素材分离。',
    '商品图作为外观参考输入，不在浏览器中抠图后叠加到背景图。为控制传输体积，商品参考图会按最大边 1536 像素缩放并转为 JPEG。',
    '引用历史效果图时，使用该结果的原始生成 URL 和元数据，不截图聊天卡片或缩略图。',
])

doc.add_heading('6.3 点选位置与尺寸', level=2)
add_bullets(doc, [
    '用户在完整空间图上单击目标位置；该点表示新增物体“底部中心接触点”。',
    '前端保存归一化坐标：x_ratio = clickX / displayedWidth，y_ratio = clickY / displayedHeight。',
    '生成前再将归一化坐标映射回原图像素，以避免页面缩放导致落点偏移。',
    '长度、宽度、高度为可选项。用户提供时优先作为比例约束；未提供时由模型根据类别与空间参考估算。',
    '当前不承诺工程测量精度；对明显放不下的尺寸做生成前提醒属于后续规划。',
])

doc.add_heading('6.4 添加与替换', level=2)
add_table(doc, ['操作', '行为', '保护规则'], [
    ['添加 ADD', '把商品参考物加入目标位置', '保留原房间结构和现有物体，只在新物体附近做必要融合'],
    ['替换 REPLACE', '移除点选附近的目标家具并放入新商品', '只删除被替换家具；合理情况下保留其上方/附近独立小物品'],
], widths=[3.0, 6.0, 7.0])

doc.add_heading('6.5 冲突检测与分级纠错', level=2)
add_table(doc, ['情形', '系统处理', '用户选项'], [
    ['点选在右侧、文字明确说左侧等左右冲突', '生成前阻断并提示不一致', '以点选位置为准 / 以文字描述为准'],
    ['落点轻微偏离可摆放面', '自动纠正，不打断用户', '无'],
    ['落地灯点在墙中部、沙发明显放不下等', '生成前提示建议位置', '使用建议位置 / 仍按原位置'],
    ['“商品完全不变”同时要求“改颜色”等语义冲突', '当前不能稳定识别', '后续规划：生成前语义冲突选择'],
], widths=[5.0, 6.2, 4.8])

doc.add_heading('6.6 结果卡片与连续编辑', level=2)
add_bullets(doc, [
    '最终效果图在聊天流中与 AI 消息一起展示。正常情况下前台展示 Round 2 最终图，不展示中间过程图。',
    '结果卡片提供“引用这张图、下载、重新生成”。引用后在输入区展示类似聊天引用的卡片，并将该图设为下一轮主编辑图。',
    '继续编辑时可保留位置、风格或商品主体，只修改用户本轮明确提出的内容。',
    '若 Round 2 已生成但 QC 服务不可用，仍展示真实 Round 2 图并标记“QC 待确认”，允许查看、引用和下载。',
])

page_break(doc)
doc.add_heading('7. 管理员后台', level=1)
add_para(doc, '管理员后台位于同一网页顶部页签，用于验证真实生成链路，而不是面向普通用户展示。当前没有账号、角色或权限隔离，因此只能视为演示后台。')

doc.add_heading('7.1 AI 生成任务', level=2)
add_bullets(doc, [
    '显示任务编号、所属对话题目、操作类型、基础图版本、状态和过程图数量。',
    '任务详情固定展示 Round 0、Round 1、Round 2、Round 3。',
    'Round 1 和 Round 2 显示两张实际中间图；Round 0 和 Round 3 只显示结构化结果，不生成图片。',
    '每轮显示 Prompt、prompt_version、last_updated_at、start_time、end_time、duration、request_id、输入和输出。',
    '首次 QC 失败时允许一次自动纠错；后台仍只保留并展示 Round 2 的最终纠错图，同时在详情中记录初次尝试与纠错尝试。',
])

doc.add_heading('7.2 Prompt 管理', level=2)
add_table(doc, ['能力', '当前行为'], [
    ['独立管理', 'Round 0、Round 1、Round 2、Round 3 可分别查看和编辑'],
    ['编辑流程', '默认查看模式；支持编辑、保存、取消、恢复默认 Prompt'],
    ['版本', '保存任一轮后生成新的 prompt_version，并更新 last_updated_at'],
    ['生效规则', '保存即生效，新任务使用最新版本'],
    ['历史审计', '历史任务保留创建时实际使用的 Prompt 快照，不随模板更新改变'],
    ['存储范围', '配置保存在当前浏览器 localStorage；服务器集中配置待确认'],
], widths=[4.0, 12.0])

doc.add_heading('8. 用户流程', level=1)
doc.add_heading('8.1 文字流程与连接关系', level=2)
add_numbered(doc, [
    '进入页面 → 选择历史对话或新建对话。',
    '上传空间原图 → 上传商品参考图。',
    '选择“添加”或“替换” → 在空间图上点选底部中心落点 → 可选填写尺寸 → 输入指令。',
    '点击发送 → 校验必填项 → 检测点选/文字冲突 → 评估落点合理性。',
    '如有明确冲突或明显不合理落点 → 用户选择优先级或是否采用建议位置 → 回到生成。',
    '创建并冻结任务归属与上下文 → 依次执行需求解析、基础生成、综合修正、质量检查。',
    '成功 → 展示最终图 → 用户下载、引用或继续调整。',
    '失败 → 展示错误与“重新生成” → 沿用上一任务素材和参数再次执行。',
])

doc.add_heading('8.2 可编辑 Mermaid 流程图', level=2)
user_flow = '''flowchart TD
    A[进入对话生图] --> B{选择对话}
    B -->|历史| C[恢复该对话素材与消息]
    B -->|新建| D[进入空状态]
    C --> E[上传或沿用空间图与商品图]
    D --> E
    E --> F[选择添加/替换]
    F --> G[点选落点并填写可选尺寸与指令]
    G --> H{必填项完整?}
    H -->|否| I[提示缺失内容]
    I --> G
    H -->|是| J{点选与文字冲突?}
    J -->|是| K[选择点选优先或文字优先]
    J -->|否| L{落点明显不合理?}
    K --> L
    L -->|是| M[选择建议位置或原位置]
    L -->|否| N[冻结任务ID与归属]
    M --> N
    N --> O[执行生成管线]
    O --> P{成功?}
    P -->|是| Q[展示最终图]
    Q --> R[下载 / 引用 / 继续编辑]
    P -->|否| S[错误提示与重新生成]
    S --> O'''
add_code(doc, user_flow)

doc.add_heading('8.3 英文流程图生图提示词', level=2)
add_callout(doc, 'Image-generation prompt (English)', 'Create a clean enterprise product flowchart for a Chinese AI home-preview web application. Use a white background, deep forest-green primary lines, muted orange highlights, rounded rectangles, clear arrows, and minimal flat vector styling. Show the complete user journey: select or create conversation; upload original room image and product reference; choose Add or Replace; click a bottom-center placement point; optionally enter dimensions and a natural-language instruction; validate required inputs; detect point-versus-text conflicts; confirm recommended placement for obviously invalid points; freeze task and conversation identifiers; run AI generation; then branch to success with View, Download, Quote and Continue Editing, or failure with a prominent Regenerate action. No decorative room illustration, no fake percentages, no code, no gradients. 16:9 landscape, presentation-ready, high legibility.', 'note')

page_break(doc)
doc.add_heading('9. 图像生成与质量检查流程', level=1)
doc.add_heading('9.1 方案概述', level=2)
add_para(doc, '当前方案不是浏览器抠图叠加。Worker 将完整空间图作为主编辑图，将商品图作为外观参考，调用百炼视觉理解和图像编辑模型完成生成。模型名为当前默认值，可由部署环境覆盖。')
add_table(doc, ['轮次', '是否生图', '当前默认模型/能力', '输出'], [
    ['Round 0 需求解析', '否', 'qwen3-vl-plus：理解空间、落点、尺寸、障碍、光线与冲突状态', 'scene-analysis.json'],
    ['Round 1 基础生成', '是', 'qwen-image-3.0：完成添加/替换，优先商品外观、目标位置、房间结构', 'round-1-base.png'],
    ['Round 2 综合修正', '是', 'qwen-image-3.0：修正比例、透视、接地、遮挡、穿模、阴影、光照、色温、背景漂移', 'round-2-composite.png'],
    ['Round 3 质量检查', '否', 'qwen3-vl-plus：比较基础图、商品图、用户要求与 Round 2', 'PASS/FAIL + 问题说明'],
], widths=[3.6, 2.1, 7.7, 2.6], font_size=8.6)
add_callout(doc, '自动纠错规则', 'Round 3 首次 FAIL 时，系统只允许一次 Round 2 局部纠错并再次质检；不无限重试。正常路径为两轮实际生图，只有 QC 失败时才可能多一次纠错生图。', 'info')

doc.add_heading('9.2 每轮关键约束', level=2)
add_bullets(doc, [
    '保持完整画幅：保留原图左右、上下边界；若接口尺寸受限，采用等比例缩放和 padding/letterbox，不直接裁切。',
    '保持背景：保留房间结构、相机角度、墙地纹理、曝光和颗粒；禁止全局磨皮、柔光、重调色。',
    '局部融合：只在新增/替换家具附近调整必要阴影、反射、遮挡和色温。',
    '商品一致性：参考商品的形状、颜色、材质、结构和风格，但拒绝商品页文字、价格、按钮和背景。',
    '空间真实性：目标点作为底部中心接触点，遵循地面/台面、透视、前后遮挡和合理比例。',
    '替换保护：只移除被替换物，附近独立小物品合理时继续保留。',
])

doc.add_heading('9.3 文字流程与连接关系', level=2)
add_para(doc, '任务创建后先执行 Round 0；解析成功后，Round 1 以完整基础图和商品参考图生成中间图 1；Round 2 同时参考 Round 1、商品图和原始背景锁定图生成中间图 2；Round 3 比较四类输入。PASS 直接完成；FAIL 进入唯一一次纠错与复检；QC 服务不可用但 Round 2 已成功时，结果以“QC 待确认”交付。')

doc.add_heading('9.4 可编辑 Mermaid 流程图', level=2)
pipeline_flow = '''flowchart LR
    A[冻结任务上下文] --> R0[Round 0 需求解析\n不生图]
    R0 --> R1[Round 1 基础生成\n中间图 1]
    R1 --> R2[Round 2 综合修正\n中间图 2]
    R2 --> R3[Round 3 质量检查\n不生图]
    R3 -->|PASS| S[前台展示最终图]
    R3 -->|FAIL 且未重试| C[唯一一次局部纠错]
    C --> R3B[再次质量检查]
    R3B -->|PASS 或最终 FAIL| S
    R3 -->|QC 服务不可用且 Round 2 已完成| U[展示 Round 2\n标记 QC 待确认]
    U --> S2[可查看 / 引用 / 下载]
    R0 -.每轮记录.-> LOG[Prompt 版本 / request_id / 起止时间 / 耗时 / 状态]
    R1 -.-> LOG
    R2 -.-> LOG
    R3 -.-> LOG'''
add_code(doc, pipeline_flow)

doc.add_heading('9.5 英文流程图生图提示词', level=2)
add_callout(doc, 'Image-generation prompt (English)', 'Design a precise technical pipeline diagram for a scene-aware furniture image-editing system. Use four main stages from left to right: Round 0 Requirement Analysis (no image), Round 1 Base Generation (intermediate image 1), Round 2 Integrated Spatial and Lighting Correction (intermediate image 2), and Round 3 Visual Quality Check (no image). From quality check, show PASS to final result, FAIL to exactly one localized correction and one re-check, and QC unavailable to deliver the Round 2 image with a “QC pending confirmation” label. Under all stages, show a shared audit log for prompt version, request ID, start time, end time, duration, and status. Minimal flat vector, white canvas, forest green and warm orange accents, rounded cards, subtle icons, no model logos, no fake percentages, 16:9 landscape.', 'note')

page_break(doc)
doc.add_heading('10. 页面与交互说明', level=1)
add_table(doc, ['页面/区域', '主要内容', '关键交互'], [
    ['顶部导航', '品牌、对话生图/管理员后台切换', '切换视图，不改变当前任务归属'],
    ['左侧历史对话', '搜索、新对话、标题、结果数、删除入口', '切换、归档、新建、删除确认'],
    ['中间聊天区', '素材图、用户消息、阶段进度、AI 消息、结果图', '点选落点、引用、下载、重新生成'],
    ['底部输入区', '文字指令、添加/替换、商品图状态、落点状态', '发送并生成'],
    ['右侧素材区', '空间原图、商品参考图、尺寸、点选坐标', '上传/更换素材、填写尺寸'],
    ['管理员后台', '任务列表、轮次详情、Prompt 管理', '审阅过程图、编辑 Prompt'],
], widths=[3.0, 6.2, 6.8])

doc.add_heading('10.1 阶段式等待体验', level=2)
add_para(doc, '前台显示阶段进度，不显示虚假的精确实时百分比：')
add_numbered(doc, ['正在理解需求', '正在生成基础效果', '正在优化空间与光影', '正在检查结果', '完成'])
add_para(doc, '用户可以切换到其他历史对话；任务完成后必须写回创建任务时冻结的 conversation_id，而不是写入完成时当前打开的对话。')

doc.add_heading('10.2 必填校验', level=2)
add_table(doc, ['缺失项', '提示/阻断'], [
    ['文字指令为空', '提示输入调整内容，不创建任务'],
    ['空间图为空', '提示上传空间图片'],
    ['商品图为空', '提示上传商品参考图'],
    ['未点选落点', '提示在完整图片范围内点选位置'],
], widths=[5.0, 11.0])

doc.add_heading('10.3 产品截图插入建议', level=2)
add_para(doc, '下列截图不是理解正文的前提，用于评审或发布时补充视觉证据。')
add_table(doc, ['建议位置', '需要截取的页面与操作', '应展示的状态', '建议裁剪范围'], [
    ['本节后', '对话生图页：上传空间图和商品图、选 ADD、在图中点选', '聊天区、点选标记、右侧尺寸与素材', '保留左侧历史、中间主图、右侧素材；裁掉浏览器标签栏'],
    ['6.6 后', '引用一张已生成效果图并输入继续修改', '引用卡片、引用来源 V1/V2、发送区域', '只保留聊天消息和输入框'],
    ['7.1 后', '管理员后台打开一个完成任务', 'Round 0–3、Round 1/2 图、耗时与 request_id', '任务详情卡完整宽度，不包含无关浏览器区域'],
    ['7.2 后', 'Prompt 管理点击某一轮“编辑”', 'prompt_version、last_updated_at、保存/取消/恢复默认', 'Prompt 卡片区域'],
    ['11.2 后', '制造明显左右冲突或不合理落点', '二选一确认弹窗', '主图与弹窗同时可见'],
    ['15 后', '测试结果表或测试工作表', 'T01–T12 与 V1/V2 结论', '仅截表格标题、编号、结论和问题类型；图片可缩略'],
], widths=[2.4, 6.3, 4.2, 3.1], font_size=8.3)

page_break(doc)
doc.add_heading('11. 异常情况与恢复策略', level=1)
add_table(doc, ['异常', '当前处理', '限制/后续'], [
    ['必填内容为空', '发送前提示并阻断', '已实现'],
    ['图片格式或大小不符合', '上传时提示；空间图限制 10MB', '商品图更细粒度限制待确认'],
    ['点选/文字左右冲突', '生成前二选一确认', '更广泛语义矛盾识别未完成'],
    ['落点轻微偏差', '自动纠正', '纠正结果不单独可视化'],
    ['落点明显不合理', '询问采用建议位置或原位置', '建议点仅坐标，无图上标记'],
    ['百炼请求失败/网络失败', '错误卡片 + 重新生成；保留本轮上下文', '刷新后若大图素材未能本地保存，直接重试可能不可用'],
    ['Round 3 QC 500，但 Round 2 成功', '展示 Round 2，标记 QC 待确认', '质量风险交由用户判断'],
    ['任务生成中切换对话', '结果写回创建时对话', '已实现冻结 task/conversation/scene/base image'],
    ['删除正在生成的对话', '禁止删除并提示等待完成', '取消任务功能未实现'],
    ['生成耗时过长', '阶段状态持续显示', '无单轮超时、队列状态和离页后台续跑提示'],
], widths=[3.8, 7.4, 4.8], font_size=8.7)

doc.add_heading('12. 数据与状态', level=1)
doc.add_heading('12.1 任务冻结字段', level=2)
add_para(doc, '点击生成时立即固定 task_id、conversation_id、scene_id、base_image_id，以及操作类型、点选位置、尺寸、文字指令、素材来源和 Prompt 快照。这样即使用户切换对话，结果也能写回正确位置。')

doc.add_heading('12.2 当前存储事实', level=2)
add_table(doc, ['数据', '当前存储', '说明'], [
    ['历史对话与消息 HTML', '浏览器 localStorage', '约保留最近 12 个对话'],
    ['任务摘要与轮次', '浏览器 localStorage', '约保留最近 20 个任务；大图源会从持久化摘要中剔除以控制容量'],
    ['Prompt 配置', '浏览器 localStorage', '保存即生效，仅当前浏览器'],
    ['生成图 URL', '百炼返回的远程 URL + 浏览器会话状态', 'URL 有效期和长期存档策略待确认'],
    ['API Key', '部署运行时环境变量', '不应写入前端或本 PRD'],
], widths=[4.0, 4.6, 7.4], font_size=8.8)
add_callout(doc, '待确认', '生产环境是否需要服务端数据库、对象存储、用户账号、跨设备同步、数据保留周期、远程文件删除和权限审计，目前材料无法确认。', 'warning')

doc.add_heading('12.3 图片元数据与坐标映射', level=2)
add_code(doc, '''original_width: 用户上传原图宽度
original_height: 用户上传原图高度
aspect_ratio = original_width / original_height
x_ratio = clickX / displayedWidth
y_ratio = clickY / displayedHeight
original_x_px = x_ratio * original_width
original_y_px = y_ratio * original_height''')

page_break(doc)
doc.add_heading('13. 功能范围', level=1)
add_table(doc, ['范围分类', '内容'], [
    ['当前已实现', '聊天式生图；空间图/商品图上传；添加/替换；点选；尺寸；冲突确认；分级纠错；两轮生图；QC；阶段进度；结果引用/下载/重试；独立历史；删除；管理员过程审阅；Prompt 管理；原图比例保护'],
    ['部分实现', '尺寸约束、复杂遮挡、小物品保留、背景锁定、自动纠错、历史版本关系。能力存在，但生成结果仍受模型与输入质量影响。'],
    ['明确未实现', '账号与权限、云端持久化、跨设备同步、多人协作、任务取消、正式版本树/分支对比、局部蒙版重做、商品链接导入、2–3 个可视化推荐位置、生产级监控与告警'],
    ['待确认', '目标商业模式、付费方式、正式 SLA、服务端数据模型、管理员角色范围、支持的正式图片格式/大小、生成图片长期保存策略'],
], widths=[3.2, 12.8])

doc.add_heading('14. 验收标准', level=1)
doc.add_heading('14.1 页面与基础交互', level=2)
add_table(doc, ['编号', '验收项', '通过条件'], [
    ['A01', 'HTML 正常打开', '页面无阻断性脚本错误；主要布局可见'],
    ['A02', '按钮可操作', '新对话、切换、上传、添加/替换、发送、引用、下载、重试、后台页签可响应'],
    ['A03', '必填校验', '缺指令、空间图、商品图或落点时明确提示且不创建任务'],
    ['A04', '图片上传', '支持的格式可预览；原始宽高和比例被记录'],
    ['A05', '历史独立', '新对话不残留上轮内容；任务完成写回创建时对话'],
    ['A06', '删除', '二次确认；进行中对话不可删；删除后正确切换'],
], widths=[1.6, 4.2, 10.2], font_size=8.8)

doc.add_heading('14.2 真实接口与等待状态', level=2)
add_table(doc, ['编号', '验收项', '通过条件'], [
    ['B01', '真实调用', '后台存在百炼 request_id；前台不使用商品叠加兜底图'],
    ['B02', '阶段进度', '依次显示理解、基础生成、空间与光影优化、检查、完成；无虚假百分比'],
    ['B03', '成功展示', '接口成功后展示真实 Round 2 最终图，可引用和下载'],
    ['B04', '失败反馈', '显示可理解的错误信息和显眼“重新生成”按钮'],
    ['B05', '重试复用', '重试沿用原 base image、商品图、操作、落点、尺寸和指令'],
    ['B06', 'QC 降级', 'Round 2 成功而 QC 不可用时，仍交付 Round 2 并标为待确认'],
], widths=[1.6, 4.2, 10.2], font_size=8.8)

doc.add_heading('14.3 图像质量', level=2)
add_table(doc, ['编号', '检查项', '通过条件'], [
    ['C01', '完整构图', '与原图保持相同构图意图，无页面截图或固定比例裁切'],
    ['C02', '商品一致性', '形状、颜色、材质、结构和风格可识别；无商品页文字/价格/按钮'],
    ['C03', '落点与比例', '家具接近用户选择位置；尺寸不明显过大或过小'],
    ['C04', '空间融合', '透视、接地、遮挡、阴影和光照合理，无明显悬空/穿模'],
    ['C05', '背景保护', '未编辑区结构、纹理、曝光、颗粒和颜色尽量不变，无全局柔光磨皮'],
    ['C06', '替换保护', '只删除被替换家具；附近独立物品合理时保留'],
    ['C07', '失败上限', 'QC FAIL 最多自动纠错一次，不无限重试'],
], widths=[1.6, 4.2, 10.2], font_size=8.8)

doc.add_heading('14.4 管理员后台', level=2)
add_table(doc, ['编号', '验收项', '通过条件'], [
    ['D01', '过程图', '每个正常任务显示 Round 1 与 Round 2 两张过程图，不显示第三张常规中间图'],
    ['D02', '轮次记录', 'Round 0–3 均含状态、Prompt、起止时间、耗时、输入/输出'],
    ['D03', 'Prompt 版本', '新任务用最新版本；历史任务显示实际快照'],
    ['D04', 'Prompt 编辑', '四轮分别支持编辑、保存、取消、恢复默认，保存即生效'],
], widths=[1.6, 4.2, 10.2], font_size=8.8)

page_break(doc)
doc.add_heading('15. 测试集与 V1/V2 结果', level=1)
add_para(doc, '测试集共 12 条，覆盖正例、连续编辑、历史引用、负例、低质量输入、复杂遮挡、点文冲突和稳定性。结论来自所提供测试表；“通过”不等于生产级稳定。')
test_rows = [
    ['T01', '正例·添加沙发', '通过', '通过', 'V1 暴露固定比例裁切；当前基线已改为原图输入与比例保护'],
    ['T02', '正例·替换沙发', '通过', '通过', '替换区域和背景保持符合预期'],
    ['T03', '正例·有尺寸比例', '失败', '通过', 'V1 百炼 503；V2 尺寸适配成功'],
    ['T04', '正例·无尺寸估算', '通过', '通过', 'V1 曾意外裁切；V2 结果可用'],
    ['T05', '多轮·主体一致', '通过', '通过', '引用上一轮缩小沙发，连续编辑可用'],
    ['T06', '历史引用加落地灯', '通过', '半通过', 'V2 落点偏近沙发；与历史裁切/位置精度有关'],
    ['T07', '负例·不合理落点', '半通过', '通过', 'V2 有纠错确认；推荐位置仍仅显示坐标'],
    ['T08', '负例·口语指令', '通过', '通过', '可理解非标准表达'],
    ['T09', '边界·暗图/低质量', '通过', '通过', '可生成，但极模糊时仍缺少重新上传提示'],
    ['T10', '边界·复杂遮挡', '半通过', '通过', 'V2 改善花瓶遮挡与前后关系'],
    ['T11', '对抗·点文冲突', '半通过', '通过', 'V2 生成前触发优先级选择'],
    ['T12', '对抗·矛盾与稳定性', '半通过', '失败', '任务等待/处理 1223.8s，未生成结果'],
]
add_table(doc, ['编号', '测试目的', 'V1', 'V2', '关键观察'], test_rows, widths=[1.2, 4.0, 1.5, 1.7, 7.6], font_size=7.9)
add_callout(doc, '汇总结论', 'V2 严格通过 10/12，半通过 1/12，失败 1/12；若将半通过视为“主任务可完成但有明显质量问题”，则 11/12 完成主流程。样本量小，结果只能用于当前版本回归与问题优先级判断。', 'info')

doc.add_heading('15.1 主要测试发现', level=2)
add_bullets(doc, [
    '添加、替换、口语化输入、无尺寸估算、连续引用和复杂遮挡在当前样例中总体可用。',
    'V2 改善了任务归属、失败重试、点文冲突、落点确认、后台可审计和背景保护。',
    '减少生图轮次后，落点/几何修正可能变弱，T06 体现为落地灯与沙发距离偏近。',
    '超长任务仍是关键稳定性问题，T12 达到 1223.8 秒并失败。',
    '低质量输入即便能生成，也缺少清晰的“建议重新上传”路径。',
])

doc.add_heading('16. V1→当前 V2 的重要修改', level=1)
add_table(doc, ['优先级', '原问题', '当前修改', '现状'], [
    ['P0', '生成结果写入完成时打开的错误对话', '创建任务时冻结 task/conversation/scene/base image', '已修复'],
    ['P0', '失败后没有显眼重试入口', '失败卡片新增重新生成并复用上下文', '已修复'],
    ['P0', '点选与文字方位冲突直接生成', '生成前阻断并提供两种优先级', '左右冲突已实现；通用语义冲突未完成'],
    ['P0', '不合理落点静默纠正', '轻微自动纠正、明显错误先确认', '已实现；建议点可视化未完成'],
    ['P1', '三轮以上生图耗时长', '压缩为两轮实际生图 + 一轮 QC', '已实现；长任务仍可能超过 1000 秒'],
    ['P1', '背景磨皮、柔光和重调色', 'Prompt 锁定纹理/曝光/颗粒，只允许局部融合', '已实现约束；模型仍可能偶发漂移'],
    ['P1', '多物体遮挡、穿模和小物品丢失', 'Round 2 与 QC 增加穿模/遮挡/悬空/保留物检查', '已增强，仍受模型能力限制'],
    ['P1', '后台 Prompt 不可编辑', '四轮独立查看/编辑/版本化', '已实现本地配置'],
    ['P1', '历史对话不能删除', '弱化删除入口、二次确认、进行中阻断', '已实现本地删除'],
    ['P1', '预览截图或固定画布导致裁切', '原始图片直送、比例保留、归一化坐标、letterbox', '当前代码已修复'],
    ['P1', 'Round 2 成功但 QC 500 导致前台无图', 'QC 不可用时交付 Round 2 并标记待确认', '已修复并兼容部分历史失败任务'],
], widths=[1.5, 4.7, 6.2, 3.6], font_size=8.0)

page_break(doc)
doc.add_heading('17. 已知问题', level=1)
add_table(doc, ['优先级', '问题', '影响', '当前建议'], [
    ['P0', '家具落点精度下降', '引用旧图继续加物时可能偏离点选位置或靠近已有家具', '强化 Round 2 对点选距离、周边障碍和接地的优先级'],
    ['P0', '建议位置不可视化', '用户只看到坐标，看不到系统建议放哪里', '确认弹窗中在原图上显示绿色点/虚线圈'],
    ['P1', '通用语义冲突未识别', '如“完全不变”与“改颜色”同时出现仍可能直接生图', '在 Round 0 前增加轻量语义冲突解析'],
    ['P1', '长任务/模型排队', '单任务可能等待 1000 秒以上，用户无法判断排队还是故障', '设置单轮超时、最大重试、排队状态与离页续跑'],
    ['P1', '极模糊或低质量素材缺少强提示', '系统可能勉强生成而非建议重传', '上传后做清晰度/曝光评估并提供重新上传'],
    ['P1', '本地存储容量与跨设备限制', '历史和重试上下文可能因容量、刷新或设备变化丢失', '服务端任务/对话存储和对象存储'],
    ['P1', '管理员无鉴权', '普通用户可见后台页签和 Prompt 配置', '生产环境增加身份、角色与审计'],
    ['P2', '生成 URL 长期有效性待确认', '历史图片可能无法长期访问', '将结果归档到受控对象存储并设置保留策略'],
], widths=[1.5, 4.3, 5.5, 4.7], font_size=8.2)

doc.add_heading('18. 后续规划', level=1)
add_table(doc, ['阶段', '计划项', '价值', '状态'], [
    ['近期 P0', '推荐落点图上可视化；加强 Round 2 点位/几何约束', '提升摆放准确性和用户理解', '未实现'],
    ['近期 P1', '语义冲突检测扩展到颜色、材质、保持不变等要求', '减少无效生图和费用', '未实现'],
    ['近期 P1', '单轮超时、队列状态、最大重试与后台续跑', '解决 1000 秒级等待体验', '未实现'],
    ['近期 P1', '图片清晰度/曝光检测和重新上传提示', '减少垃圾输入导致的低质量结果', '未实现'],
    ['中期', '尺寸适配提醒：明显放不下时生成前提示', '降低比例错误', '未实现'],
    ['中期', '局部重做/蒙版，只重做阴影或局部区域', '减少背景漂移和耗时', '未实现'],
    ['中期', '给出 2–3 个可视化推荐摆放位置', '降低点选成本', '未实现'],
    ['中期', '服务端历史、图片归档、账号和权限', '支持跨设备、审计和生产使用', '待产品与架构确认'],
    ['远期', '商品链接导入，自动读取商品图与尺寸', '减少手工上传和输入', '未实现'],
], widths=[2.2, 6.4, 4.8, 2.6], font_size=8.3)
add_callout(doc, '不应误写为规划', '“历史对话自动生成标题”已在当前代码中实现，因此已归入现有能力，不再列为后续功能。', 'note')

doc.add_heading('19. 待确认事项', level=1)
add_table(doc, ['主题', '待确认问题'], [
    ['业务', '正式目标用户优先级、商业模式、计费方式、是否接入电商链路'],
    ['模型', '生产环境固定模型、区域、配额、并发限制、超时与降级策略'],
    ['数据', '是否建设服务端 Conversation/Scene/ImageVersion/Task 数据模型及保留周期'],
    ['安全', '管理员身份、Prompt 编辑权限、API Key 管理、日志脱敏、用户图片隐私策略'],
    ['质量', '各家具品类的可接受落点误差、尺寸误差和背景漂移阈值'],
    ['运营', '是否需要人工审核、失败任务工单、成本报表和模型切换'],
    ['兼容性', '正式支持的浏览器、移动端布局、图片格式/大小和无障碍要求'],
], widths=[3.0, 13.0])

page_break(doc)
doc.add_heading('20. 产品截图与评审材料制作规范', level=1)
add_para(doc, '为避免“截图好看但不能证明功能”的问题，所有截图都应带明确任务、输入和状态。建议优先使用真实测试任务，不手工拼接结果。')
add_table(doc, ['截图编号', '操作步骤', '需要证明的事实', '裁剪/标注'], [
    ['S01 主流程', '新建对话→上传两图→点选→ADD→发送', '输入与点选都在同一任务中', '保留三栏和输入区；标注点选点'],
    ['S02 冲突确认', '点右侧，文字写左侧', '生成前出现二选一', '保留原图、点位、弹窗'],
    ['S03 结果卡', '等待成功', '真实图、引用、下载、重新生成', '保留结果图和操作区'],
    ['S04 引用续改', '点击引用并要求缩小', '历史图成为主编辑图', '保留引用卡片与本轮指令'],
    ['S05 失败恢复', '模拟接口失败', '错误信息与直接重试', '保留阶段进度、错误卡和按钮'],
    ['S06 后台过程', '打开最新成功任务', 'Round 1/2 两张图、Round 0/3 无图、耗时', '保留完整轮次卡片'],
    ['S07 Prompt 管理', '编辑 Round 2 Prompt', '版本、更新时间与编辑动作', '聚焦 Prompt 卡片'],
], widths=[2.4, 5.5, 5.5, 2.6], font_size=8.3)
add_para(doc, '建议截图统一隐藏浏览器个人账号、插件浮层、API Key 和任何用户隐私；不要在文档中展示真实密钥。', color=RED)

doc.add_heading('21. 其他配图建议', level=1)
add_para(doc, '本 PRD 不建议额外生成装饰性 AI 配图。产品逻辑更适合用真实截图、文字流程和 Mermaid 图表达；装饰图可能让评审误以为是已实现页面。')

doc.add_heading('附录 A：核心 Prompt 约束摘要', level=1)
add_para(doc, '以下为产品层摘要，不替代后台实际 Prompt：')
add_bullets(doc, [
    'preserve original room appearance; preserve original wall and floor texture; preserve exposure and image grain',
    'preserve camera angle, perspective, room structure, unedited objects and the complete frame',
    'do not beautify or smooth the background; do not apply global soft light or global color grading',
    'use the product image as appearance reference; do not copy product-page text, price, buttons or background',
    'place the furniture at the normalized bottom-center contact point with realistic scale, perspective, grounding, shadow and occlusion',
    'for replacement, remove only the selected furniture and preserve nearby independent objects when physically reasonable',
    'check object penetration, incorrect occlusion, floating objects, missing preserved objects and background drift',
])

doc.add_heading('附录 B：术语', level=1)
add_table(doc, ['术语', '说明'], [
    ['基础图 / Base image', '本轮编辑的主图，可以是用户原始空间图，也可以是引用的历史效果图'],
    ['商品参考图', '用于参考家具外观的图片，不作为浏览器前景叠加层'],
    ['落点', '用户点选的物体底部中心接触点'],
    ['Scene-aware editing', '基于空间、透视、遮挡和光影理解的局部图像编辑'],
    ['Round 1/2', '两轮实际生图：基础生成与综合修正'],
    ['QC', '质量检查；不生成图片，只判断 PASS/FAIL 和问题'],
    ['背景漂移', '未编辑区域被意外改色、磨皮、改结构或重绘'],
    ['Prompt 快照', '任务创建时实际使用的四轮 Prompt 版本副本'],
], widths=[4.0, 12.0])

doc.add_heading('附录 C：版本与发布信息', level=1)
add_table(doc, ['项', '值'], [
    ['当前基线目录', '/Users/nd/Desktop/Jiazhuang_AI/V2_test'],
    ['当前基线说明', 'V2_test/README.md、DEPLOYMENT_INFO.md、SHA256SUMS.txt'],
    ['部署项目 ID', 'appgprj_6ac1d507ee648191acf8cc8b326860eb'],
    ['冻结时网站版本', '21'],
    ['当前演示 URL', 'https://qijing-ai-home-preview.mawellhihi.chatgpt.site/'],
    ['说明', 'URL 和部署信息反映冻结时状态；后续发布可能变化，变更时应同步更新本文档。'],
], widths=[4.2, 11.8])

# Keep table rows from splitting where possible and add document metadata
for table in doc.tables:
    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        cant_split = OxmlElement('w:cantSplit')
        tr_pr.append(cant_split)

props = doc.core_properties
props.title = 'AI 家居空间预览 产品需求文档 PRD V2.1'
props.subject = '当前实现、测试结果、V1→V2 修改与后续规划'
props.author = 'Codex · 基于项目真实材料整理'
props.keywords = 'AI家居, 场景编辑, PRD, V2, 百炼, 图像生成'
props.comments = '以 2026-10-09 V2_test 基线为准；未确认项已标注。'

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
