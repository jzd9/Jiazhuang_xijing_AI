from pathlib import Path
from shutil import copy2
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

ROOT = Path('/Users/nd/Desktop/Jiazhuang_AI')
TMP = Path('/var/folders/4x/j5cwm8ds2nz4xx4jd0trzxvh0000gn/T')
OUT = ROOT / 'PRD_V2.1_AI家居空间预览_图文完整版.docx'
doc = Document(ROOT / 'PRD_V2.1_AI家居空间预览_当前实现与测试版.docx')
assets = ROOT / 'prd_assets'
assets.mkdir(exist_ok=True)
files = [
('ecab1108-bb13-41bb-8a4d-4aa31d6efef7','user_flow'),
('38110676-2303-46c3-bbd4-37f734a81410','new_conversation'),
('b0d2e20b-3279-4715-a655-eaa8d1e17616','prompt_management'),
('8eef267b-46f0-4436-ab30-88c30d1b42e0','position_conflict'),
('ee17c6fe-b3b8-424b-b47e-7478b9cf322b','tests_T01_T05'),
('161b43da-a8e4-4828-bd85-5fccd97a969e','tests_T06_T10'),
('b83f3250-85e1-41e8-b5b3-66de586be1d8','tests_T11_T12'),
('10f98f0f-5b0c-46b2-a618-94e4e14f5020','result_reference'),
('1bcee78a-b3e7-404b-8539-8080adca35b2','admin_rounds'),
('a4332e83-30e3-4b8d-aa20-e4fbbd6e8988','generation_pipeline')]
for ident, name in files:
    copy2(TMP / f'codex-clipboard-{ident}.png', assets / f'{name}.png')

def text(el):
    return ''.join(el.itertext()) if False else ''.join(el.xpath('.//w:t/text()'))

# Remove production instructions, retaining all product requirements and Mermaid sources.
remove = False
for el in list(doc.element.body):
    value = text(el)
    if any(value.startswith(s) for s in ['8.3 英文', '9.5 英文', '10.3 产品截图', '20. 产品截图', '21. 其他配图']):
        remove = True
    elif remove and any(value.startswith(s) for s in ['9. 图像', '10. 页面', '11. 异常', '附录 A']):
        remove = False
    if remove and el.tag != qn('w:sectPr'):
        doc.element.body.remove(el)

for p in doc.paragraphs:
    if '超长任务仍是关键稳定性问题' in p.text:
        p.text = '超长任务仍是关键稳定性问题。T12 已生成图片，但未先确认矛盾要求，且耗时 1223.8 秒，测试结论为不通过。'
for t in doc.tables:
    for row in t.rows:
        for cell in row.cells:
            if '任务等待/处理 1223.8s，未生成结果' in cell.text:
                cell.text = '已生成图片，但未先确认矛盾要求；耗时 1223.8s，测试不通过'
    # Synchronize column grid with existing deliberate cell widths.
    if t.rows:
        for col, cell in zip(t.columns, t.rows[0].cells):
            if cell.width:
                col.width = cell.width

def before(prefix, name, caption, note='', width=16):
    anchor = next(p for p in doc.paragraphs if p.text.startswith(prefix))
    p = anchor.insert_paragraph_before()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    shape = p.add_run().add_picture(str(assets / f'{name}.png'), width=Cm(width))
    shape._inline.docPr.set('descr', caption)
    cp = anchor.insert_paragraph_before(caption)
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.paragraph_format.space_after = Pt(8)
    for r in cp.runs: r.font.size = Pt(9)
    if note:
        np = anchor.insert_paragraph_before(note)
        for r in np.runs: r.font.size = Pt(9)

before('6.2 ', 'new_conversation', '图 1 新建对话空状态与素材设置')
before('6.6 ', 'position_conflict', '图 2 点选位置与文字描述冲突时的确认', width=13)
before('7. 管理员', 'result_reference', '图 3 结果卡片与引用效果图状态',
       '引用将选定效果图作为下一次编辑的基础图，不自动改用最新图片；截图中的 V1 是图片版本，不是软件版本。', width=14)
before('7.2 ', 'admin_rounds', '图 4 管理员任务详情与两轮实际生成结果',
       '该任务展示 Round 0 解析、Round 1 基础生成、Round 2 综合修正及 Round 3 QC，同时保留版本、请求编号、起止时间和耗时。', width=15)
before('8. 用户', 'prompt_management', '图 5 Prompt 管理查看与编辑状态')
before('8.2 ', 'user_flow', '图 6 用户操作流程示意',
       '读图说明：图中“引用用于新对话”应理解为引用到后续编辑，不代表自动新建对话；必填项缺失后须补充并重新校验。任务身份在任务创建时固定，确认弹窗仅调整本轮执行参数。')
before('9.4 ', 'generation_pipeline', '图 7 两轮生图与质量检查流程示意',
       '读图说明：Round 2 仅修正目标物体与必要局部融合，不进行全局改色或重新设计。QC 服务不可用是独立异常分支，不等同于质量 FAIL；已有 Round 2 图可标注 QC 待确认后展示。图中“合规性”不代表已实现独立内容合规系统。')
before('16. ', 'tests_T01_T05', '图 8 测试记录 T01 至 T05')
before('16. ', 'tests_T06_T10', '图 9 测试记录 T06 至 T10')
before('16. ', 'tests_T11_T12', '图 10 测试记录 T11 至 T12',
       'T12 有生成结果，但矛盾要求未先确认且耗时 1223.8 秒，仍判为不通过。测试表中的原因推测不视为已经验证的根因。')

# Remove inherited decorative title rules and excessive forced page gaps.
for style in doc.styles:
    if style.name in ('Title', 'Subtitle') or style.name.startswith('Heading'):
        style.font.color.rgb = RGBColor(0, 0, 0)
    for border in list(style.element.xpath('.//w:pBdr')):
        border.getparent().remove(border)
break_seen = False
number = 0
for p in doc.paragraphs:
    if p.style.name in ('Title', 'Subtitle') or p.style.name.startswith('Heading'):
        for r in p.runs: r.font.color.rgb = RGBColor(0, 0, 0)
    for border in list(p._p.xpath('.//w:pBdr')):
        border.getparent().remove(border)
    for br in list(p._p.xpath('.//w:br[@w:type="page"]')):
        if break_seen: br.getparent().remove(br)
        break_seen = True
    if p.style.name == 'List Number':
        number += 1
        p.style = doc.styles['Normal']
        p.text = f'{number}. {p.text}'
    else:
        number = 0
doc.save(OUT)
print(OUT)
print('images', len(doc.inline_shapes))
