from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

ROOT = Path(__file__).resolve().parents[1]
import sys as _sys
_args = _sys.argv[1:]
SOURCE = ROOT / (_args[0] if _args else '技术说明_V2_test_当前实现.md')
OUT = ROOT / (_args[1] if len(_args) > 1 else SOURCE.stem + '.docx')
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin = sec.bottom_margin = Inches(.7)
sec.left_margin = sec.right_margin = Inches(.75)

for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3']:
    style = doc.styles[name]
    style.font.name = 'Arial'
    style._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'PingFang SC')
    style.font.color.rgb = RGBColor(0,0,0)
    style.paragraph_format.space_after = Pt(7)
normal = doc.styles['Normal']
normal.font.size = Pt(11)
normal.paragraph_format.line_spacing = 1.22
normal.paragraph_format.widow_control = True
for name, size in [('Title',24),('Heading 1',17),('Heading 2',13),('Heading 3',12)]:
    s=doc.styles[name]; s.font.size=Pt(size)
    s.paragraph_format.space_before=Pt(15 if name != 'Title' else 0)
    s.paragraph_format.keep_with_next=True

def para(text, style=None, code=False):
    p=doc.add_paragraph(style=style)
    r=p.add_run(text.replace('`',''))
    if code:
        r.font.name='Menlo'; r.font.size=Pt(9)
        p.paragraph_format.line_spacing=1.05
        p.paragraph_format.space_after=Pt(1)
    return p

def table(lines):
    data=[[v.strip() for v in l.strip().strip('|').split('|')] for l in lines]
    data=[r for r in data if not all(re.fullmatch(r'[:\- ]+',v) for v in r)]
    n=len(data[0]); t=doc.add_table(rows=0,cols=n);t.autofit=False
    t.alignment=WD_TABLE_ALIGNMENT.CENTER
    fractions={2:[.32,.68],3:[.25,.39,.36]}.get(n,[1/n]*n)
    widths=[Inches(7*f) for f in fractions]
    for col,w in zip(t.columns,widths):col.width=w
    for ri,row in enumerate(data):
        cells=t.add_row().cells
        cant=OxmlElement('w:cantSplit');t.rows[ri]._tr.get_or_add_trPr().append(cant)
        if ri==0:
            h=OxmlElement('w:tblHeader');t.rows[ri]._tr.get_or_add_trPr().append(h)
        for ci,(cell,txt) in enumerate(zip(cells,row)):
            cell.width=widths[ci];cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(4);p.paragraph_format.space_before=Pt(4);p.paragraph_format.line_spacing=1.12
            r=p.add_run(txt.replace('`',''));r.font.size=Pt(10);r.bold=ri==0
            r.font.color.rgb=RGBColor.from_string('FFFFFF' if ri==0 else '222222')
            pr=cell._tc.get_or_add_tcPr()
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'243D50' if ri==0 else ('F2F5F7' if ri%2==0 else 'FFFFFF'));pr.append(shade)
            borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
            pr.append(borders)
            margins=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                e=OxmlElement('w:'+side);e.set(qn('w:w'),'80');e.set(qn('w:type'),'dxa');margins.append(e)
            pr.append(margins)
    doc.add_paragraph().paragraph_format.space_after=Pt(1)

def png_px(path):
    try:
        with open(path,'rb') as f:data=f.read(33)
        if data[:8]!=b'\x89PNG\r\n\x1a\n':return None
        return int.from_bytes(data[16:20],'big'),int.from_bytes(data[20:24],'big')
    except OSError:return None

def picture(caption, rel):
    p = ROOT / rel
    par = doc.add_paragraph(); par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    par.paragraph_format.space_before = Pt(8)
    if p.exists():
        w = min(Inches(6.9), Inches((png_px(p) or (960, 540))[0] / 96))
        par.add_run().add_picture(str(p), width=w)
    else:
        r = par.add_run(f'【图片缺失：{rel}】'); r.font.size = Pt(10)
    c = doc.add_paragraph(); c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c.paragraph_format.space_after = Pt(10)
    r = c.add_run(caption.replace('`','')); r.font.size = Pt(9); r.font.color.rgb = RGBColor.from_string('666666')

lines=SOURCE.read_text().splitlines();i=0;incode=False
while i<len(lines):
    line=lines[i]
    if line.startswith('```'):
        incode=not incode
        if incode:para('可编辑 Mermaid 源码', 'Heading 3')
        i+=1;continue
    if not line.strip():i+=1;continue
    if incode:para(line,code=True)
    elif line.startswith('|'):
        group=[]
        while i<len(lines) and lines[i].startswith('|'):group.append(lines[i]);i+=1
        table(group);continue
    elif line.startswith('# '):para(line[2:],'Title')
    elif line.startswith('## '):para(line[3:],'Heading 1')
    elif line.startswith('### '):para(line[4:],'Heading 2')
    elif re.match(r'^!\[.*?\]\(.+?\)\s*$', line):
        m = re.match(r'^!\[(.*?)\]\((.+?)\)\s*$', line)
        picture(m.group(1), m.group(2))
    else:para(line)
    i+=1

footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=footer.add_run('栖境 AI 家居空间预览  |  ');r.font.size=Pt(9)
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
doc.core_properties.title='栖境 AI 家居空间预览'
doc.core_properties.subject=SOURCE.stem
doc.core_properties.author=''
for root in (doc.styles.element, doc.element):
    for border in list(root.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
doc.save(OUT)
print(OUT)
